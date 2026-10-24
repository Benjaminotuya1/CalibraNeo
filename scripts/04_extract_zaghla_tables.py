import os
import docx
from pypdf import PdfReader
import pandas as pd

raw_dir = os.path.expanduser("~/CalibraNeo_Project/data_raw/zaghla_2026")
docx_path = os.path.join(raw_dir, "jitc-14-4-s001.docx")
pdf_path = os.path.join(raw_dir, "jitc-14-4-s002.pdf")
xlsb_path = os.path.join(raw_dir, "jitc-14-4-s003.xlsb")

print("=== 1. CHECKING TABLES INSIDE WORD DOC (s001.docx) ===")
doc = docx.Document(docx_path)
print(f"Number of tables in s001.docx: {len(doc.tables)}")
for idx, table in enumerate(doc.tables):
    print(f"  Table {idx+1}: {len(table.rows)} rows x {len(table.columns)} cols")
    if len(table.rows) > 0:
        headers = [cell.text.strip().replace("\n", " ") for cell in table.rows[0].cells]
        print(f"  Headers: {headers[:8]}")

print("\n=== 2. CHECKING PAGES INSIDE PDF (s002.pdf) ===")
reader = PdfReader(pdf_path)
print(f"Total Pages in s002.pdf: {len(reader.pages)}")
for i in range(min(5, len(reader.pages))):
    txt = reader.pages[i].extract_text()
    first_lines = [line.strip() for line in txt.split("\n") if line.strip()][:4]
    print(f"  Page {i+1} preview: {first_lines}")

print("\n=== 3. SUMMARIZING THE 403,224 ROWS IN s003.xlsb ===")
df = pd.read_excel(xlsb_path, sheet_name="CLean_Filtered Data", engine="pyxlsb")
print("Unique OriginalFileNames (Experiments):", df["OriginalFileName"].nunique())
print("Sample OriginalFileNames:")
for fname in df["OriginalFileName"].dropna().unique()[:10]:
    print("  -", fname)
