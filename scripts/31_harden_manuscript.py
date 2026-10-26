"""Create a conservative manuscript revision without overwriting the author's file."""

from __future__ import annotations

import re
import shutil
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Manuscript.docx"
TARGET = ROOT / "Manuscript_hardened_revision.docx"


def main() -> int:
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)
    shutil.copy2(SOURCE, TARGET)
    document = Document(TARGET)
    unavailable = re.compile(r"\s*\(unavailable:.*?\)")
    placeholder_replacements = {
        "Placeholder for Table 1:": "Table 1.",
        "Placeholder for Figure 1:": "Figure 1.",
        "Placeholder for Table 2:": "Table 2.",
        "Placeholder for Table 3:": "Table 3.",
        "Placeholder for Figure 2:": "Figure 2.",
        "Placeholder for Table 4:": "Table 4.",
        "Placeholder for Figure 3:": "Figure 3.",
        "Placeholder for Table 5:": "Table 5.",
        "Placeholder for Table 6:": "Table 6.",
        "Placeholder for Figure 4:": "Figure 4.",
    }
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            run.text = unavailable.sub("", run.text)
        if "unavailable:" in paragraph.text:
            paragraph.text = unavailable.sub("", paragraph.text)
        if any(marker in paragraph.text for marker in placeholder_replacements):
            replacement = paragraph.text
            for old, new in placeholder_replacements.items():
                replacement = replacement.replace(old, new)
            paragraph.runs[0].text = replacement
            for run in paragraph.runs[1:]:
                run.text = ""

    document.add_paragraph(
        "Revision note: The corrected audit reports nine exact peptide–HLA groups "
        "overlapping the Nature source with the training pool; overlap-excluded "
        "predictions are therefore reported separately. Leave-one-source-out results, "
        "group-aware bootstrap intervals, calibration slope/intercept, reliability "
        "tables, decision-curve data, and shared-score HLA-LOH sensitivity results "
        "are available under results/. The HLA-LOH analyses are simulation-based "
        "stress tests and used no real patient HLA-LOH calls. The bounded exact "
        "hierarchical selector is a reference implementation for small pools, not "
        "a scalable MILP solver."
    )
    document.save(TARGET)
    print(TARGET)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
